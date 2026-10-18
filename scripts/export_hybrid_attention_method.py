"""Export the native editable method schematic at its journal print dimensions."""
from pathlib import Path
import base64
import hashlib
import json
import re
import subprocess

from PIL import Image
from pypdf import PdfReader, PdfWriter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

OUT = Path('paper/figures/hybrid_attention_method')
BASE = 'DiffuPrimer_hybrid_attention_method'
NATIVE_PDF = OUT / 'render_native' / f'{BASE}.pdf'

reader = PdfReader(NATIVE_PDF)
page = reader.pages[0]
scale = (178 / 25.4 * 72) / float(page.mediabox.width)
page.scale_by(scale)
writer = PdfWriter()
writer.add_page(page)
writer.add_metadata({
    '/Title': 'DiffuPrimer: conditional attention for primer-pair diffusion',
    '/Subject': 'Method schematic; publication size 178 x 106.8 mm',
})
with (OUT / f'{BASE}.pdf').open('wb') as handle:
    writer.write(handle)

def rasterize(name, *resolution_args):
    subprocess.run([
        'pdftoppm', '-png', '-singlefile', *resolution_args,
        str(OUT / f'{BASE}.pdf'), str(OUT / name),
    ], check=True)

rasterize(BASE, '-r', '350')
rasterize('preview', '-scale-to-x', '1800', '-scale-to-y', '-1')
rasterize('print_review_700px', '-scale-to-x', '700', '-scale-to-y', '-1')

with Image.open(OUT / f'{BASE}.png') as source:
    image = source.convert('RGB')
    image.save(OUT / f'{BASE}.png', dpi=(350, 350))
    image.save(OUT / f'{BASE}.tiff', compression='tiff_lzw', dpi=(350, 350))
    pixels = image.size

# Embed the regular/bold open font so that the standalone vector SVG is portable.
svg_path = OUT / f'{BASE}.svg'
svg = svg_path.read_text()
svg = re.sub(r'<style(?: id="embedded-dejavu-fonts")?>.*?</style>', '', svg, flags=re.S)
styles = []
for filename, weight in [('DejaVuSans.ttf', '400'), ('DejaVuSans-Bold.ttf', '700')]:
    font_path = Path('/usr/share/fonts/truetype/dejavu') / filename
    encoded = base64.b64encode(font_path.read_bytes()).decode()
    styles.append(
        '@font-face{font-family:"DejaVu Sans";font-weight:' + weight +
        ';src:url(data:font/ttf;base64,' + encoded + ') format("truetype");}'
    )
svg = svg.replace('<defs>', '<defs><style id="embedded-dejavu-fonts">' + ''.join(styles) + '</style>', 1)
svg_path.write_text(svg)

presentation = Presentation(OUT / f'{BASE}.pptx')
shapes = list(presentation.slides[0].shapes)
font_sizes = []
for shape in shapes:
    if not shape.has_text_frame:
        continue
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            size = run.font.size or paragraph.font.size
            if run.text and size is not None:
                font_sizes.append(size.pt)
assert min(font_sizes) >= 16
assert not any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes)
assert len(presentation.slides) == 1
assert abs(float(page.mediabox.width) / 72 * 25.4 - 178) < 1e-6

report = {
    'figure_role': 'Scientific method schematic, without experimental result claims',
    'architecture': 'Candidate hybrid architecture: 4 residual dilated local CNN blocks (96 channels, dilations1/4/16/64), then2 conditional Transformer blocks (4 heads, rotary positions, FiLM). Validation pending.',
    'editable_source': f'{BASE}.pptx',
    'figure_generator': 'scripts/draw_hybrid_attention_method.py',
    'export_script': 'scripts/export_hybrid_attention_method.py',
    'reference_style': 'https://doi.org/10.1093/bib/bbaf308',
    'native_editable_objects': len(shapes),
    'embedded_raster_objects_in_ppt': 0,
    'publication_width_mm': 178,
    'publication_height_mm': 106.8,
    'minimum_native_font_pt': min(font_sizes),
    'minimum_print_font_pt': round(min(font_sizes) * scale, 3),
    'caption_native_font_pt': 17,
    'caption_print_font_pt': round(17 * scale, 3),
    'raster_dpi': 350,
    'raster_pixels': pixels,
    'review_preview_width_px': 700,
    'svg_fonts_embedded': True,
    'files': {
        path.name: {
            'bytes': path.stat().st_size,
            'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
        }
        for path in OUT.iterdir()
        if path.is_file() and path.suffix in {'.pptx', '.pdf', '.svg', '.png', '.tiff'}
    },
}
(OUT / 'provenance.json').write_text(json.dumps(report, indent=2) + '\n')
print(json.dumps({key: value for key, value in report.items() if key != 'files'}, indent=2))
