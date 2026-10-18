"""Editable PowerPoint and vector SVG method figure, from shared primitives."""
from pathlib import Path
import math, html
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from lxml import etree
OUT=Path('paper/figures/cnn_publication_method'); OUT.mkdir(parents=True,exist_ok=True)
prs=Presentation(); prs.slide_width= Inches(18); prs.slide_height=Inches(16.0)
s=prs.slides.add_slide(prs.slide_layouts[6]); svg=[]
C={'ink':'193D50','muted':'526B79','teal':'278E76','coral':'D97554','gold':'B78929','purple':'7F6FA9','gray':'DCE5E9','line':'CEDADF'}
def rect(x,y,w,h,fill='FFFFFF',stroke=None,r=.1):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
 sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(fill); sh.line.fill.background() if not stroke else None
 if stroke: sh.line.color.rgb=RGBColor.from_string(stroke); sh.line.width=Pt(.8)
 if r:
  try:sh.adjustments[0]=.08
  except Exception:pass
 svg.append(f'<rect x="{x*100}" y="{y*100}" width="{w*100}" height="{h*100}" rx="{r*100}" fill="#{fill}" stroke="#{stroke or fill}" stroke-width="1"/>')
 return sh
def text(x,y,w,h,t,size=16,color='193D50',bold=False,align='left'):
 sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 from pptx.enum.text import PP_ALIGN
 for i,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(0)
  p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
  xx=x*100+(w*50 if align=='center' else w*100 if align=='right' else 0)
  svg.append(f'<text x="{xx}" y="{y*100+size*100/72*.9+i*size*100/72*1.2}" font-family="DejaVu Sans" font-size="{size*100/72}" font-weight="{700 if bold else 400}" text-anchor="'+{'left':'start','center':'middle','right':'end'}[align]+f'" fill="#{color}">{html.escape(line)}</text>')
 return sh
def line(x1,y1,x2,y2,color='607887',width=1.5,arrow=False,dash=False):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2));sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(width)
 if arrow:
  el=etree.SubElement(sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln'),'{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd');el.set('type','triangle');el.set('w','med');el.set('len','med')
 svg.append(f'<line x1="{x1*100}" y1="{y1*100}" x2="{x2*100}" y2="{y2*100}" stroke="#{color}" stroke-width="{width*100/72}"'+(' marker-end="url(#arrow)"' if arrow else '')+'/>')
def cells(x,y,states,cell=.19):
 colors={'0':C['gray'],'F':C['teal'],'R':C['coral'],'M':C['purple']}
 for i,k in enumerate(states):
  rect(x+i*cell,y,cell-.025,.28,colors[k],r=.025)
  if cell>=.18:text(x+i*cell,y+.035,cell-.025,.2,k,9,'FFFFFF' if k!='0' else C['muted'],True,'center')
def pair(x,y,w=2.4,f=.35,r=1.8,muted=False):
 line(x,y,x+w,y,'B8C8CE',1.6); line(x+f,y-.10,x+f+.4,y-.1,C['teal'],3,True);line(x+r+.4,y+.1,x+r,y+.1,C['coral'],3,True)
def badge(x,y,n,title,w):
 rect(x,y,.36,.36,C['ink'],r=.08);text(x,y+.035,.36,.28,n,15,'FFFFFF',True,'center');text(x+.5,y+.015,w-.5,.45,title,19,C['ink'],True)
# Biological illustration: editable ribbons, nucleotide beads and primer glyphs.
def poly(points,fill):
 b=s.shapes.build_freeform(points[0][0]*1000,points[0][1]*1000,scale= Inches(1)/1000)
 b.add_line_segments([(x*1000,y*1000) for x,y in points[1:]],close=True)
 sh=b.convert_to_shape();sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background()
 svg.append('<polygon points="'+' '.join(f'{x*100},{y*100}' for x,y in points)+f'" fill="#{fill}"/>')
 return sh
def disk(x,y,r,fill,stroke=None):
 sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x-r), Inches(y-r), Inches(r*2), Inches(r*2));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
 if stroke:sh.line.color.rgb=RGBColor.from_string(stroke);sh.line.width=Pt(.8)
 else:sh.line.fill.background()
 svg.append(f'<circle cx="{x*100}" cy="{y*100}" r="{r*100}" fill="#{fill}" stroke="#{stroke or fill}"/>')
def ribbon(points,width,fill):
 upper=[];lower=[]
 for i,(x,y) in enumerate(points):
  a=points[max(0,i-1)];b=points[min(len(points)-1,i+1)];dx=b[0]-a[0];dy=b[1]-a[1];n=math.hypot(dx,dy) or 1
  upper.append((x-dy/n*width/2,y+dx/n*width/2));lower.append((x+dy/n*width/2,y-dx/n*width/2))
 poly(upper+lower[::-1],fill)
def helix(x,y,w,amp=.53,cycles=2.1,pale=False):
 col1,col2=('BFDAD8','D3CFE4') if pale else ('54A99E','9382B9')
 n=400
 for j in range(0,n+1,14):
  xx=x+w*j/n;a=amp*math.sin(j/n*math.pi*2*cycles)
  line(xx,y-a,xx,y+a,'E3DCCA' if pale else 'D5C5A0',1.9)
 for sign,col in [(1,col1),(-1,col2)]:
  pts=[(x+w*i/n,y+sign*amp*math.sin(i/n*math.pi*2*cycles)) for i in range(n+1)]
  ribbon(pts,.07 if pale else .065,col)
  # a narrower highlight gives the editable strand a soft ribbon appearance
  ribbon([(a,b-.016) for a,b in pts],.016,'E7F1EF' if pale else 'BDDDD7')
def arrow_curve(points,color,width=.08,head=.23):
 ribbon(points,width,color)
 x,y=points[-1];a,b=points[-2];dx=x-a;dy=y-b;n=math.hypot(dx,dy);ux=dx/n;uy=dy/n
 poly([(x+ux*head*.45,y+uy*head*.45),(x-ux*head+uy*head*.6,y-uy*head-ux*head*.6),(x-ux*head-uy*head*.6,y-uy*head+ux*head*.6)],color)
def ntrow(x,y,seq,step=.23,size=13):
 colors={'A':'347C71','T':'AA6637','C':'416F92','G':'7B5999'}
 for i,nt in enumerate(seq):
  disk(x+i*step+.1,y+.11,.105,'FFFFFF')
  text(x+i*step-.025,y-.025,.25,.34,nt,size,colors.get(nt,'6C8590'),True,'center')
def labels(x,y,states,step=.23):
 for i,k in enumerate(states):
  col={'F':C['teal'],'R':C['coral'],'M':'A395C4','0':'B7C8CC'}[k]
  if k=='0':disk(x+i*step+.10,y+.13,.037,col)
  else:
   disk(x+i*step+.10,y+.13,.107,col)
   text(x+i*step,y+.007,.20,.27,k,10,'FFFFFF',True,'center')
def primer(x,y,w,col,right=True):
 line(x,y,x+w,y,col,4)
 # tapered primer extension, not an administrative flowchart arrow
 end=x+w if right else x;dx=1 if right else -1
 poly([(end+dx*.12,y),(end-dx*.11,y-.09),(end-dx*.11,y+.09)],col)
def lock(x,y,col):
 pts=[(x+.075+.075*math.cos(math.pi-i*math.pi/20),y+.08-.09*math.sin(math.pi-i*math.pi/20)) for i in range(21)]
 ribbon(pts,.028,col);rect(x-.01,y+.07,.17,.14,col,r=.025);disk(x+.075,y+.125,.018,'FFFFFF')
def duplex(x,y,w,ff=.3,rr=2.4,shade=False):
 if shade:rect(x+ff,y-.015,rr-ff+.48,.20,'F1EBD9',r=.025)
 line(x,y,x+w,y,'BDCDCE',1.2);line(x,y+.17,x+w,y+.17,'BDCDCE',1.2)
 for i in range(int(w/.13)+1):line(x+i*.13,y+.03,x+i*.13,y+.14,'DFE5E4',.6)
 primer(x+ff,y+.13,.44,C['teal'],True);primer(x+rr,y+.035,.44,C['coral'],False)

def textline(x,y,w,t,size=17,color=None,bold=False,align='left'):
 return text(x,y,w,.42,t,size,color or C['ink'],bold,align)

def state_row(x,y,states):
 # Large, native editable characters; unselected positions remain explicit zeros.
 for i,k in enumerate(states):
  col={'F':C['teal'],'R':C['coral'],'M':C['purple'],'0':'6C8590'}[k]
  if k=='0':
   text(x+i*.27-.025,y,.275,.35,'0',17,col,False,'center')
  else:
   disk(x+i*.27+.11,y+.165,.129,col)
   text(x+i*.27-.025,y+.005,.27,.35,k,16,'FFFFFF',True,'center')

def biological_state(x,heading,symbol,states):
 textline(x,1.41,4.78,heading,20,bold=True,align='center')
 textline(x,1.92,4.78,symbol,18,C['purple'],align='center')
 line(x-.04,2.9,x+4.69,2.9,'FFFFFF',24)
 ntrow(x,2.73,'ATGCGTACGATCGACTGA',.27,17)
 state_row(x,3.59,states)
 for i,k in enumerate(states):
  if k in 'FR':line(x+i*.27+.11,3.03,x+i*.27+.11,3.46,'B5D8CD' if k=='F' else 'E6BDAF',1.3)

# Print typography: 19 pt on the drawing canvas = 7.40 pt at 178 mm.
def tx(x,y,w,t,size=19,col=None,bold=False,align='left',h=.65):
 return text(x,y,w,h,t,size,col or C['ink'],bold,align)
def arr(x,y,xx,yy,col=None):line(x,y,xx,yy,col or C['muted'],1.7,True)
def band(y,h):rect(.35,y,17.3,h,'F6F9F6',r=.1)
def heading(y,k,title):
 disk(.65,y+.17,.19,C['ink']);tx(.43,y-.015,.44,k,19,'FFFFFF',True,'center');tx(.99,y-.035,16.3,title,22,bold=True)
rect(0,0,18,16.0,r=0)
tx(.47,.16,17.1,'DiffuPrimer | Conditional diffusion of primer-binding positions',24,bold=True)
heading(.83,'A','Learn to recover a primer pair on a fixed transcript')
# Faint native double helix is biological context, not generated sequence.
helix(.65,2.00,16.7,.36,4.5,True)
seq='ATGCGTACGATCGACTGA'
for x,title,states in [(.70,'Clean map  x₀','00FFF000000RRR0000'),(6.45,'Partial mask  xₜ','0MFMF0MM0MMRMR0M0M'),(12.20,'Initialization  xT','M'*18)]:
 tx(x,1.35,4.92,title,20,bold=True,align='center')
 line(x,2.03,x+4.8,2.03,'FFFFFF',23)
 ntrow(x,1.87,seq,.27,19)
 for i,k in enumerate(states):
  col={'0':'6C8590','F':C['teal'],'R':C['coral'],'M':C['purple']}[k]
  tx(x+i*.27-.015,2.41,.28,k,19,col,k!='0','center')
for x in [5.70,11.45]:arr(x,1.72,x+.42,1.72,C['purple'])
tx(5.85,2.87,6.3,'Training: mask with probability t / T',19,C['purple'],align='center')
arrow_curve([(15.80-13.1*i/100,3.34+.19*math.sin(math.pi*i/100)) for i in range(101)],C['teal'],.047,.14)
tx(5.65,3.15,6.9,'Reverse: conditional recovery, T → 1',19,C['teal'],True,'center')
for x,key,meaning,col in [(.75,'0','background','6C8590'),(4.50,'F','forward',C['teal']),(7.65,'R','reverse',C['coral']),(10.85,'M','mask',C['purple'])]:
 tx(x,3.77,.35,key,19,col,True);tx(x+.4,3.77,3,meaning,19,C['muted'])
tx(13.20,3.77,4.1,'Template unchanged',19,C['muted'],align='right')

# Architecture receives the largest share of the figure.
band(4.35,9.38)
heading(4.62,'B','A conditional dilated CNN guides compatible paired-interval sampling')
# Explicit representations and dimensionality.
tx(.72,5.16,4.1,'Sequence + noisy map',20,bold=True)
ntrow(.78,5.73,'ATCGN',.35,19)
tx(2.68,5.70,2.0,'L × 96',19,C['muted'])
tx(.76,6.13,4.1,'DNA: 5 → 96 embedding',19)
tx(.76,6.52,4.1,'0/F/R/M: 4 → 96',19)

tx(5.16,5.16,5.15,'Design + diffusion time',20,bold=True)
tx(5.16,5.71,5.15,'c: lengths, Tm, GC (9 values)',19)
tx(5.16,6.10,5.3,'[τ, τ², sin(πτ), cos(πτ)]',19)
tx(5.16,6.49,5.3,'Linear 13 → 96 → 96; SiLU',19)

tx(11.18,5.16,5.9,'Relative transcript position',20,bold=True)
tx(11.18,5.71,5.9,'[i / L, sin(2πi / L)]',19)
tx(11.18,6.10,5.9,'1 × 1 Conv: 2 → 96',19)
tx(11.18,6.49,5.9,'τ = t / T; L = valid length',19,C['muted'])
# Summation is faithful to input injection, without repeated FiLM claims.
for xx in [2.7,8.4,14.0]:line(xx,7.00,xx,7.17,'92A8A4',1.3)
line(2.7,7.17,14,7.17,'92A8A4',1.3)
arr(8.4,7.17,8.4,7.43,'92A8A4')
disk(8.4,7.61,.20,'FFFFFF',C['teal']);tx(8.19,7.40,.42,'+',22,C['teal'],True,'center')
tx(9.01,7.39,7.95,'Add; broadcast condition over L → 96 × L',19,C['muted'])
# 4 native CNN blocks: actual dilated taps and true residual additions.
xs=[.84,5.05,9.26,13.47]
for j,(xx,d,rf) in enumerate(zip(xs,[1,4,16,64],[5,21,85,341])):
 rect(xx,8.10,3.62,1.97,'FFFFFF','C6D8CF',.12)
 tx(xx+.1,8.24,3.42,f'Block {j+1}  ·  d = {d}',20,bold=True,align='center')
 # Same 5-tap kernel, wider spacing = increased dilation (schematic scale).
 spacing=[.20,.30,.41,.52][j];mid=xx+1.81
 taps=[mid+(q-2)*spacing for q in range(5)]
 line(xx+.30,9.29,xx+3.32,9.29,'C9D8D3',1.2)
 for z in range(17):disk(xx+.31+z*.188,9.29,.027,'BBCBC6')
 for tap in taps:
  line(tap,9.25,mid,8.96,'88B3A6',1.1);disk(tap,9.29,.065,C['teal'])
 disk(mid,8.91,.07,C['purple'])
 tx(xx+.1,9.59,3.42,f'Conv RF = {rf} nt',19,C['muted'],align='center')
 # Identity bypass explicitly feeds the addition at block output.
 line(xx+.15,8.99,xx+.15,7.88,'76948C',1.1)
 line(xx+.15,7.88,xx+3.48,7.88,'76948C',1.1)
 arr(xx+3.48,7.88,xx+3.48,8.89,'76948C')
 disk(xx+3.48,9.05,.12,'FFFFFF',C['teal']);tx(xx+3.34,8.90,.28,'+',19,C['teal'],True,'center')
 if j<3:arr(xx+3.67,9.05,xx+4.10,9.05)
line(8.17,7.61,.57,7.61,C['teal'],1.3)
line(.57,7.61,.57,9.05,C['teal'],1.3)
arr(.57,9.05,.81,9.05,C['teal'])
tx(.80,10.27,16.5,'Each block: GroupNorm → SiLU → Conv1d (k = 5) → SiLU → 1 × 1 Conv + residual',19)
tx(.80,10.70,16.5,'96 channels; length preserved; valid-site masking. Conv RF: convolutional span; GN: GroupNorm.',19,C['muted'])
# Prediction head and structured recovery occupy a full clear row.
line(17.22,9.05,17.40,9.05,C['teal'],1.2)
line(17.40,9.05,17.40,11.05,C['teal'],1.2)
line(17.40,11.05,3.42,11.05,C['teal'],1.2)
arr(3.42,11.05,3.42,11.22,C['teal'])
rect(.80,11.24,5.25,.60,'E7F0EC',r=.08)
tx(.98,11.32,4.90,'Head: GN → SiLU → 1 × 1 Conv',19)
arr(6.13,11.54,6.60,11.54)
tx(6.82,11.26,3.3,'L × 3 logits',20,bold=True)
tx(6.82,11.65,3.3,'0 / F / R at each site',19,C['muted'])
arr(10.17,11.54,10.64,11.54)
tx(10.84,11.25,6.45,'Sample one compatible F/R interval pair',19,bold=True)
tx(10.84,11.66,6.45,'Lengths · order · anchors · eligibility · ΔTm',19,C['muted'])
# Natural recovery line, extraction, and subsequent assessment.
tx(.78,12.25,16.5,'Train against clean x₀: masked-site cross-entropy × T / t, averaged over all valid positions.',19)
tx(.78,12.71,16.5,'Reveal M sites with probability 1 / t; preserve known labels. Iterate CNN + paired sampling to t = 1.',19,C['teal'],True)
tx(.78,13.17,16.5,'Extract F; reverse-complement R. Then assess heterodimers and transcriptome/genome specificity.',19)
heading(14.01,'C','Four biological design modes')
for xx,title in zip([.75,5.02,9.29,13.56],['Diverse pairs','Length control','Fix either primer','Local redesign']):
 tx(xx,14.59,3.80,title,20,bold=True)
for yy,f,r in [(15.21,.16,2.45),(15.55,.39,2.18)]:duplex(.83,yy,3.65,f,r,True)
for yy,r in [(15.21,1.60),(15.55,2.56)]:duplex(5.10,yy,3.65,.20,r,True)
for yy,side in [(15.21,'F'),(15.55,'R')]:
 duplex(9.37,yy,3.65,.22,2.57,True)
 lock(9.72 if side=='F' else 12.03,yy-.18,C['teal'] if side=='F' else C['coral'])
for yy,r in [(15.21,1.82),(15.55,2.57)]:
 duplex(13.64,yy,3.65,.22,r,True);lock(13.99,yy-.18,C['teal'])
# All artwork and characters are native Office objects; no image layers.
for shape in s.shapes:
 for el in shape._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):el.set('idx','0')
 etree.SubElement(shape._element.spPr,'{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
# Keep the PPT itself at publication size, scaling geometry and text together.
scale=178/25.4/18
for shape in s.shapes:
 shape.left=round(shape.left*scale);shape.top=round(shape.top*scale)
 shape.width=round(shape.width*scale);shape.height=round(shape.height*scale)
 if shape.has_text_frame:
  for p in shape.text_frame.paragraphs:
   if p.font.size:p.font.size=Pt(p.font.size.pt*scale)
   for r in p.runs:
    if r.font.size:r.font.size=Pt(r.font.size.pt*scale)
 if shape.line.width:shape.line.width=round(shape.line.width*scale)
prs.slide_width=Inches(18*scale);prs.slide_height=Inches(16.0*scale)
prs.save(OUT/'DiffuPrimer_cnn_method.pptx')
head='<svg xmlns="http://www.w3.org/2000/svg" width="178mm" height="158.2222mm" viewBox="0 0 1800 1600"><defs><marker id="arrow" viewBox="0 0 10 10" refX="8" refY="5" markerWidth="5" markerHeight="5" orient="auto-start-reverse"><path d="M 0 0 L 10 5 L 0 10 z" fill="context-stroke"/></marker></defs>'
(OUT/'DiffuPrimer_cnn_method.svg').write_text(head+''.join(svg)+'</svg>')
print(f'Saved {len(s.shapes)} editable objects at 178 mm.')
