"""Native scientific architecture + binding-map modeling figure, newly composed."""
from pathlib import Path
import html
from pptx import Presentation
from pptx.util import Mm,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from lxml import etree
O=Path('paper/figures/architecture_method');O.mkdir(parents=True,exist_ok=True)
P=Presentation();P.slide_width=Mm(178);P.slide_height=Mm(116)
S=P.slides.add_slide(P.slide_layouts[6]);V=[]
I='263E49';G='94B2A2';F='2C8F7A';R='D67B61';M='A093BC';BLUE='DCE8F4';GREEN='E3EEDD';PEACH='F6E4DA';PURPLE='EAE4F1';GRAY='DCE3E5';MUTED='617580'
def box(x,y,w,h,fill='FFFFFF',stroke=None,r=.8):
 a=S.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Mm(x),Mm(y),Mm(w),Mm(h));a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string(fill)
 if stroke:a.line.color.rgb=RGBColor.from_string(stroke);a.line.width=Pt(.55)
 else:a.line.fill.background()
 if r:a.adjustments[0]=min(.13,r/min(w,h))
 V.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{r}" fill="#{fill}" stroke="#{stroke or fill}" stroke-width=".19"/>');return a

def txt(x,y,w,t,size=7,bold=False,color=I,align='center'):
 a=S.shapes.add_textbox(Mm(x),Mm(y),Mm(w),Mm(9));f=a.text_frame;f.clear();f.word_wrap=False;f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
 for j,tline in enumerate(t.split('\n')):
  p=f.paragraphs[0] if j==0 else f.add_paragraph();p.text=tline;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(0);p.space_before=Pt(0);p.alignment={'center':PP_ALIGN.CENTER,'left':PP_ALIGN.LEFT,'right':PP_ALIGN.RIGHT}[align]
  xx=x+w/2 if align=='center' else x+w if align=='right' else x
  V.append(f'<text x="{xx}" y="{y+.315*size+j*.417*size}" font-family="DejaVu Sans" font-size="{size*.35278}" font-weight="{700 if bold else 400}" fill="#{color}" text-anchor="'+{'center':'middle','left':'start','right':'end'}[align]+f'">{html.escape(tline)}</text>')
 return a

def wire(x,y,xx,yy,arrow=True,color=I,dash=False,width=.6):
 a=S.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Mm(x),Mm(y),Mm(xx),Mm(yy));a.line.color.rgb=RGBColor.from_string(color);a.line.width=Pt(width)
 if dash:
  from pptx.enum.dml import MSO_LINE_DASH_STYLE
  a.line.dash_style=MSO_LINE_DASH_STYLE.DASH
 if arrow:
  e=etree.SubElement(a._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln'),'{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd');e.set('type','triangle');e.set('w','sm');e.set('len','sm')
 V.append(f'<line x1="{x}" y1="{y}" x2="{xx}" y2="{yy}" stroke="#{color}" stroke-width="{width*.35278}"'+(' stroke-dasharray="1,.8"' if dash else '')+(' marker-end="url(#arr)"' if arrow else '')+'/>')
def op(x,y,w,h,t,fill=GREEN):box(x,y,w,h,fill,'B2C3BB');txt(x,y+.6,w,t)
def plus(x,y):
 a=S.shapes.add_shape(MSO_SHAPE.OVAL,Mm(x-1.55),Mm(y-1.55),Mm(3.1),Mm(3.1));a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string('FFFFFF');a.line.color.rgb=RGBColor.from_string(I);a.line.width=Pt(.65)
 V.append(f'<circle cx="{x}" cy="{y}" r="1.55" fill="white" stroke="#{I}" stroke-width=".23"/>')
 wire(x-.9,y,x+.9,y,False);wire(x,y-.9,x,y+.9,False)
def gap_tip(x,y):
 a=S.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,Mm(x-.45),Mm(y),Mm(.9),Mm(.6));a.rotation=180;a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string(I);a.line.fill.background()
 V.append(f'<polygon points="{x-.45},{y} {x+.45},{y} {x},{y+.6}" fill="#{I}"/>')
def tensor(x,y,w=12,h=6,cols=6,rows=4):
 box(x+1.0,y-.9,w,h,'EDF2F6','BACBD4',0)
 for r in range(rows):
  for c in range(cols):box(x+c*w/cols,y+r*h/rows,w/cols-.12,h/rows-.12,['C7D9E8','DFEAF1','AAC6DB'][(r+c)%3],None,0)
def states(x,y,seq,step=3.5):
 for j,z in enumerate(seq):
  co={'0':GRAY,'F':F,'R':R,'M':M}[z];box(x+j*step,y,step-.13,3.8,co,None,.15);txt(x+j*step,y+.45,step-.13,z,bold=z!='0',color='FFFFFF' if z!='0' else MUTED)
def pair(x,y,w,fs=3,rs=None,anchor=None):
 rs=w-8 if rs is None else rs
 box(x+fs,y-.3,rs-fs+5,2.1,'F5ECD7',None,0)
 wire(x,y,x+w,y,False,'AABEB7',width=.5);wire(x,y+1.5,x+w,y+1.5,False,'AABEB7',width=.5)
 wire(x+fs,y+1.5,x+fs+5,y+1.5,True,F,width=1.5);wire(x+rs+5,y,x+rs,y,True,R,width=1.5)
 if anchor:
  ax=x+fs+1 if anchor=='F' else x+rs+2
  box(ax,y-2.2,1.6,1.8,'FFFFFF',F if anchor=='F' else R,.7)
  box(ax-.3,y-1.3,2.2,1.6,F if anchor=='F' else R,None,.25)
# Main division: computational network at left, biological state model at right.
box(1,1,96,88,'FFFFFF',G,1.8);box(99,1,78,88,'FFFFFF',G,1.8)
txt(4,3,90,'a  Conditional CNN architecture',9.2,True,align='left')
txt(102,3,72,'b  Diffusion of binding positions',9.2,True,align='left')
# Macro architecture: two actual embeddings plus two conditional side branches.
op(5,13,19,4.7,'DNA A',BLUE);op(27,13,19,4.7,'Noisy map xₜ',PURPLE)
wire(14.5,18,14.5,20);wire(36.5,18,36.5,20)
op(5,20.4,19,5.3,'Embed 5 → 96',BLUE);op(27,20.4,19,5.3,'Embed 4 → 96',PURPLE)
wire(14.5,26,14.5,30,False);wire(36.5,26,36.5,30,False);wire(14.5,30,36.5,30,False);wire(25.5,30,25.5,34.3)
txt(52,12,41,'Position: 2 features',7,bold=True)
op(53,16.5,39,5.2,'1 × 1 Conv: 2 → 96',BLUE)
wire(53,19.1,49,19.1,False,color='719ABB');wire(49,19.1,49,32,False,color='719ABB');wire(49,32,27,35.3,color='719ABB')
txt(52,25,41,'Design c (9) + time (4)',7,bold=True)
op(53,29.5,39,5.2,'MLP: 13 → 96 → 96',PEACH)
wire(53,32.1,51,32.1,False,color='BC967D');wire(51,32.1,51,40.5,False,color='BC967D');wire(51,40.5,27,36.9,color='BC967D')
plus(25.5,36);txt(4,34,16,'Add',7,color=MUTED)
wire(25.5,37.7,25.5,42)
tensor(19.5,43,12,5.8);txt(4,44,14,'96 × L',7,color=MUTED)
wire(25.5,49.3,25.5,51.2)
op(9,51.7,33,10,'Residual CNN × 4\nd = 1, 4, 16, 64',GREEN)
wire(25.5,62.0,25.5,65)
op(9,65.5,33,8.8,'GN → SiLU\n1 × 1 Conv: 96 → 3',BLUE)
wire(25.5,74.5,25.5,77.4)
tensor(20.5,78.3,10,5.1,cols=3,rows=4)
txt(4,84.8,43,'L × 3 logits: 0 / F / R',7,bold=True)
# True macro-to-micro zoom: five operators and the full identity branch.
wire(42,52.2,54,44.5,False,color='8A9B9D',dash=True);wire(42,61.2,54,83.5,False,color='8A9B9D',dash=True)
box(54,43.4,40,40.3,'F8FAF8','BACBC1',2)
txt(55,44.7,38,'One residual block',7.3,True)
for yy,tt,co in [(49.3,'GroupNorm',BLUE),(54.2,'SiLU',PEACH),(59.1,'Conv1d · k = 5, d',GREEN),(64.0,'SiLU',PEACH),(68.9,'Conv1d · k = 1',GREEN)]:
 op(59,yy,27,3.7,tt,co)
 if yy>49.3:gap_tip(72.5,yy-.9)
wire(72.5,48.2,72.5,49.1)
wire(72.5,48.2,90.5,48.2,False);wire(90.5,48.2,90.5,77,False);wire(90.5,77,74.2,77)
wire(72.5,72.9,72.5,75.3);plus(72.5,77);wire(72.5,78.7,72.5,81.7)
txt(56,85,37,'Width 96; length preserved',7,color=MUTED)
# Biological representation: immutable bases, clean and masked labels aligned.
txt(103,12,70,'Fixed reference transcript',7.4,True)
seq='ATGCTACGATCGACTG'
for j,base in enumerate(seq):txt(112+j*3.5,17,3.4,base,7.3,True,color=I)
wire(110,21.6,170,21.6,False,'BBCFC7',width=.55)
txt(101,25.1,10,'x₀',7.4,True);states(112,24.5,'00FFF0000RRR0000')
txt(101,39,10,'xₜ',7.4,True);states(112,38.3,'M0FMFMMM0RMRM0MM')
wire(117.3,29,117.3,37.6,color=M,width=.9)
txt(119,31.1,23,'Train: mask t / T',7,color='847199',align='left')
wire(165,37.5,165,29,color=F,width=.9)
txt(142,31.1,22,'Reverse recovery',7,color=F)
txt(106,44,66,'0 background   F forward   R reverse   M mask',7,color=MUTED)
# Joint sampling is a concrete reverse transition, not a separate downstream filter.
box(103,49.3,70,24.4,'F6F9F5','B9CCBF',1.3)
txt(105,50.6,66,'Reverse recovery: t = T → 1',7.5,True)
op(105,55.7,14,8,'CNN\nscores',BLUE)
op(123,55.7,30,8,'Compatible\nF / R spans',GREEN)
op(157,55.7,14,8,'Reveal\n1 / t',PEACH)
wire(119.4,59.7,122.5,59.7);wire(153.4,59.7,156.5,59.7)
wire(164,64,164,68,False,color=F);wire(164,68,112,68,False,color=F);wire(112,68,112,64.2,color=F)
txt(106,69.8,65,'Eligibility · ΔTm · anchors; keep known labels',7,color=MUTED)
txt(104,76,69,'Extract F; reverse-complement R',7.3,bold=True)
pair(111,81,55,5,41)
txt(102,85.7,73,'Then: heterodimer + full-reference assessment',7,color=MUTED)
# Lower strip: four biological scenes, aligned and uncluttered.
wire(2,92,176,92,False,G,width=.65)
txt(4,93.4,170,'c  One model, four design modes',9.1,True,align='left')
for x,title in [(5,'Diverse primer pairs'),(49,'Product-length control'),(93,'Fixed-primer completion'),(137,'Local redesign')]:txt(x,99,37,title,7.3,True)
for y,fs,rs in [(106,2,24),(111,6,27)]:pair(6,y,34,fs,rs)
for y,rs in [(106,18),(111,27)]:pair(50,y,34,2,rs)
pair(94,106,34,2,26,'F');pair(94,111,34,2,26,'R')
pair(138,106,34,2,18,'F');pair(138,111,34,2,26,'F')
for sh in S.shapes:
 for e in sh._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):e.set('idx','0')
 etree.SubElement(sh._element.spPr,'{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
P.save(O/'DiffuPrimer_method.pptx')
head='<svg xmlns="http://www.w3.org/2000/svg" width="178mm" height="116mm" viewBox="0 0 178 116"><defs><marker id="arr" viewBox="0 0 10 10" refX="8" refY="5" markerWidth="4" markerHeight="4" orient="auto-start-reverse"><path d="M0 0 L10 5 L0 10Z" fill="context-stroke"/></marker></defs>'
(O/'DiffuPrimer_method.svg').write_text(head+''.join(V)+'</svg>')
print(len(S.shapes),'native objects')
