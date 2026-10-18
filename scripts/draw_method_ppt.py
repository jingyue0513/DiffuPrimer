"""Editable PowerPoint and vector SVG method figure, from shared primitives."""
from pathlib import Path
import math, html
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from lxml import etree
OUT=Path('paper/figures/method_ppt'); OUT.mkdir(parents=True,exist_ok=True)
prs=Presentation(); prs.slide_width= Inches(18); prs.slide_height=Inches(11.3)
s=prs.slides.add_slide(prs.slide_layouts[6]); svg=[]
C={'ink':'193D50','muted':'607887','teal':'168D86','coral':'D97554','gold':'B78929','purple':'897CBA','gray':'DCE5E9','line':'CEDADF'}
def rect(x,y,w,h,fill='FFFFFF',stroke=None,r=.1):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
 sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(fill); sh.line.fill.background() if not stroke else None
 if stroke: sh.line.color.rgb=RGBColor.from_string(stroke); sh.line.width=Pt(.8)
 if r:
  try:sh.adjustments[0]=.08
  except Exception:pass
 svg.append(f'<rect x="{x*100}" y="{y*100}" width="{w*100}" height="{h*100}" rx="{r*100}" fill="#{fill}" stroke="#{stroke or fill}" stroke-width="1"/>')
 return sh
def text(x,y,w,h,t,size=16,color='193D50',bold=False,align='left'):
 sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 from pptx.enum.text import PP_ALIGN
 for i,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(0)
  p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
  xx=x*100+(w*50 if align=='center' else w*100 if align=='right' else 0)
  svg.append(f'<text x="{xx}" y="{y*100+size*100/72*.9+i*size*100/72*1.2}" font-family="DejaVu Sans" font-size="{size*100/72}" font-weight="{700 if bold else 400}" text-anchor="'+{'left':'start','center':'middle','right':'end'}[align]+f'" fill="#{color}">{html.escape(line)}</text>')
 return sh
def line(x1,y1,x2,y2,color='607887',width=1.5,arrow=False,dash=False):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2));sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(width)
 if arrow:
  el=etree.SubElement(sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln'),'{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd');el.set('type','triangle');el.set('w','med');el.set('len','med')
 svg.append(f'<line x1="{x1*100}" y1="{y1*100}" x2="{x2*100}" y2="{y2*100}" stroke="#{color}" stroke-width="{width*100/72}"'+(' marker-end="url(#arrow)"' if arrow else '')+'/>')
def cells(x,y,states,cell=.19):
 colors={'0':C['gray'],'F':C['teal'],'R':C['coral'],'M':C['purple']}
 for i,k in enumerate(states):
  rect(x+i*cell,y,cell-.025,.28,colors[k],r=.025)
  if cell>=.18:text(x+i*cell,y+.035,cell-.025,.2,k,9,'FFFFFF' if k!='0' else C['muted'],True,'center')
def pair(x,y,w=2.4,f=.35,r=1.8,muted=False):
 line(x,y,x+w,y,'B8C8CE',1.6); line(x+f,y-.10,x+f+.4,y-.1,C['teal'],3,True);line(x+r+.4,y+.1,x+r,y+.1,C['coral'],3,True)
def badge(x,y,n,title,w):
 rect(x,y,.36,.36,C['ink'],r=.08);text(x,y+.035,.36,.28,n,15,'FFFFFF',True,'center');text(x+.5,y+.015,w-.5,.45,title,19,C['ink'],True)
# Biological illustration: editable ribbons, nucleotide beads and primer glyphs.
def poly(points,fill):
 b=s.shapes.build_freeform(points[0][0]*1000,points[0][1]*1000,scale= Inches(1)/1000)
 b.add_line_segments([(x*1000,y*1000) for x,y in points[1:]],close=True)
 sh=b.convert_to_shape();sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background()
 svg.append('<polygon points="'+' '.join(f'{x*100},{y*100}' for x,y in points)+f'" fill="#{fill}"/>')
 return sh
def disk(x,y,r,fill,stroke=None):
 sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x-r), Inches(y-r), Inches(r*2), Inches(r*2));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
 if stroke:sh.line.color.rgb=RGBColor.from_string(stroke);sh.line.width=Pt(.8)
 else:sh.line.fill.background()
 svg.append(f'<circle cx="{x*100}" cy="{y*100}" r="{r*100}" fill="#{fill}" stroke="#{stroke or fill}"/>')
def ribbon(points,width,fill):
 upper=[];lower=[]
 for i,(x,y) in enumerate(points):
  a=points[max(0,i-1)];b=points[min(len(points)-1,i+1)];dx=b[0]-a[0];dy=b[1]-a[1];n=math.hypot(dx,dy) or 1
  upper.append((x-dy/n*width/2,y+dx/n*width/2));lower.append((x+dy/n*width/2,y-dx/n*width/2))
 poly(upper+lower[::-1],fill)
def helix(x,y,w,amp=.53,cycles=2.1,pale=False):
 col1,col2=('BFDAD8','D3CFE4') if pale else ('54A99E','9382B9')
 n=400
 for j in range(0,n+1,14):
  xx=x+w*j/n;a=amp*math.sin(j/n*math.pi*2*cycles)
  line(xx,y-a,xx,y+a,'E3DCCA' if pale else 'D5C5A0',1.9)
 for sign,col in [(1,col1),(-1,col2)]:
  pts=[(x+w*i/n,y+sign*amp*math.sin(i/n*math.pi*2*cycles)) for i in range(n+1)]
  ribbon(pts,.07 if pale else .065,col)
  # a narrower highlight gives the editable strand a soft ribbon appearance
  ribbon([(a,b-.016) for a,b in pts],.016,'E7F1EF' if pale else 'BDDDD7')
def arrow_curve(points,color,width=.08,head=.23):
 ribbon(points,width,color)
 x,y=points[-1];a,b=points[-2];dx=x-a;dy=y-b;n=math.hypot(dx,dy);ux=dx/n;uy=dy/n
 poly([(x+ux*head*.45,y+uy*head*.45),(x-ux*head+uy*head*.6,y-uy*head-ux*head*.6),(x-ux*head-uy*head*.6,y-uy*head+ux*head*.6)],color)
def ntrow(x,y,seq,step=.23,size=13):
 colors={'A':'448E84','T':'C98653','C':'6596BC','G':'A181B4'}
 for i,nt in enumerate(seq):
  disk(x+i*step+.1,y+.11,.105,'FFFFFF')
  text(x+i*step,y-.012,.20,.27,nt,size,colors[nt],True,'center')
def labels(x,y,states,step=.23):
 for i,k in enumerate(states):
  col={'F':C['teal'],'R':C['coral'],'M':'A395C4','0':'B7C8CC'}[k]
  if k=='0':disk(x+i*step+.10,y+.13,.037,col)
  else:
   disk(x+i*step+.10,y+.13,.107,col)
   text(x+i*step,y+.007,.20,.27,k,10,'FFFFFF',True,'center')
def primer(x,y,w,col,right=True):
 line(x,y,x+w,y,col,4)
 # tapered primer extension, not an administrative flowchart arrow
 end=x+w if right else x;dx=1 if right else -1
 poly([(end+dx*.12,y),(end-dx*.11,y-.09),(end-dx*.11,y+.09)],col)
def lock(x,y,col):
 pts=[(x+.075+.075*math.cos(math.pi-i*math.pi/20),y+.08-.09*math.sin(math.pi-i*math.pi/20)) for i in range(21)]
 ribbon(pts,.028,col);rect(x-.01,y+.07,.17,.14,col,r=.025);disk(x+.075,y+.125,.018,'FFFFFF')
def duplex(x,y,w,ff=.3,rr=2.4,shade=False):
 if shade:rect(x+ff,y-.015,rr-ff+.48,.20,'F1EBD9',r=.025)
 line(x,y,x+w,y,'BDCDCE',1.2);line(x,y+.17,x+w,y+.17,'BDCDCE',1.2)
 for i in range(int(w/.13)+1):line(x+i*.13,y+.03,x+i*.13,y+.14,'DFE5E4',.6)
 primer(x+ff,y+.13,.44,C['teal'],True);primer(x+rr,y+.035,.44,C['coral'],False)
# white scientific canvas and typographic hierarchy
rect(0,0,18,11.3,r=0)
text(.65,.28,5.6,.55,'DiffuPrimer',30,C['ink'],True)
text(5.17,.43,12,.37,'Conditional discrete diffusion for flexible qPCR primer design',18,C['muted'])
# faint biological background spans the entire modeling scene
helix(.4,3.10,17.2,.72,5.0,True)
text(.7,1.11,4.6,.31,'BINDING-POSITION MODELING',12,C['muted'],True)
text(6.08,1.04,7.9,.35,'Forward process  ·  progressively mask labels',15,'8B79AD')
line(5.35,1.28,5.91,1.28,'9B8AB9',1.5,True)
seq='ATGCGTACGATCGACTGA'; clean='00FFF000000RRR0000';partial='0MFMF0MM0MMRMR0M0M'
for x,head,sub,states in [(.91,'Clean binding map','x₀',clean),(6.95,'Partially masked','xₜ',partial),(12.99,'Fully masked','xT','M'*18)]:
 text(x,1.78,4.2,.37,head,19,C['ink'],True,'center')
 text(x,2.29,4.2,.32,sub,14,'8B79AD',False,'center')
 # short fixed transcript strand stays visible during the label process
 line(x-.04,3.075,x+4.07,3.075,'FFFFFF',17)
 ntrow(x,2.95,seq)
 labels(x,3.75,states)
 # tiny alignment guides connect sequence and selected/unknown positions
 for i,k in enumerate(states):
  if k in 'FR':line(x+i*.23+.1,3.23,x+i*.23+.1,3.65,'BCD9D4' if k=='F' else 'EACDBE',1)
text(.85,3.33,4.4,.25,'Reference transcript A',11,C['muted'],False,'center')
text(6.9,3.33,4.4,.25,'Same reference sequence',11,C['muted'],False,'center')
text(12.9,3.33,4.4,.25,'Same reference sequence',11,C['muted'],False,'center')
# paired opposite directions tell the corruption/reconstruction story
for xa,xb in [(5.45,6.46),(11.48,12.51)]:
 arrow_curve([(xa,2.67),(xa+.25,2.60),(xb-.25,2.60),(xb,2.67)],'A797C6',.045,.14)
 # organic masking particles
 for j in range(5):disk(xa+.16+j*.145,3.81+.08*math.sin(j*2),.025+j%2*.009,'C8BDDC')
# reveal/recovery path beneath the label states, right-to-left
pts=[(15.07-12.02*i/100,4.20+.83*math.sin(math.pi*i/100)) for i in range(101)]
arrow_curve(pts,'379E92',.075,.23)
text(6.4,4.39,5.3,.35,'Conditional denoising',19,C['teal'],True,'center')
text(6.0,5.15,6.1,.31,'Recover F / R positions, one reveal step at a time',13,C['muted'],False,'center')
# clean primer mapping motif directly below the recovered state
primer(1.5,4.59,.63,C['teal'],True);primer(3.81,4.59,.63,C['coral'],False)
text(1.32,4.93,3.43,.33,'Extract F and reverse-complement R',10.5,C['muted'],False,'center')
# Small explicit legend keeps symbolic masking clear, without a panel
for x,k,lab,col in [(13.0,'F','Forward',C['teal']),(15.02,'R','Reverse',C['coral']),(13.0,'M','Masked','A395C4')]:
 yy=4.67 if k!='M' else 5.08;disk(x+.09,yy+.12,.09,col);text(x,yy+.01,.18,.23,k,9,'FFFFFF',True,'center');text(x+.28,yy,1.5,.28,lab,11,C['muted'])
disk(15.11,5.20,.035,'B7C8CC');text(15.3,5.08,1.6,.28,'Background',11,C['muted'])
# Mechanism below: open biological/network glyphs, no module boxes.
text(1.15,5.75,4.4,.32,'Design conditions',16,C['ink'],True)
# length ruler glyph
line(1.23,6.42,2.19,6.42,C['gold'],2)
for i in range(6):line(1.23+i*.192,6.35,1.23+i*.192,6.5,C['gold'],1)
text(2.45,6.26,3.4,.55,'Length · Tm · GC\nOptional fixed F or R',12,C['muted'])
# restrained neural mesh glyph of the conditional denoiser
nodes=[[(6.15,6.15),(6.15,6.47),(6.15,6.79)],[(6.64,5.99),(6.64,6.31),(6.64,6.63),(6.64,6.95)],[(7.13,6.15),(7.13,6.47),(7.13,6.79)]]
for a,b in zip(nodes,nodes[1:]):
 for ax,ay in a:
  for bx,by in b:line(ax,ay,bx,by,'D5CBDF',.65)
for group in nodes:
 for nx,ny in group:disk(nx,ny,.065,'A294BC')
text(7.4,5.95,3.2,.32,'Conditional CNN denoiser',13,C['ink'],True)
text(7.4,6.43,3.25,.48,'Template + labels + c + t\nPredict 0 / F / R',11,C['muted'])
arrow_curve([(5.25,6.42),(5.45,6.42),(5.64,6.42)],'B6A57A',.035,.12)
# pair-guidance glyph and connection to denoising arc
text(11.99,5.85,5.0,.32,'Compatible pair sampling',15,C['ink'],True)
duplex(12.07,6.45,3.65,.18,2.76,True)
text(12.02,6.93,5,.35,'Length · order · anchors · joint ΔTm',11,C['muted'])
arrow_curve([(10.87,6.44),(11.12,6.44),(11.41,6.44)],'B6A57A',.035,.12)
# subtle upward arrows visually link mechanism to reverse diffusion
arrow_curve([(12.04,5.62),(11.79,5.27),(11.85,4.89)],'A3CFC6',.035,.12)
text(.75,7.62,14,.38,'ONE MODEL, MULTIPLE DESIGN MODES',14,C['muted'],True)
line(.75,8.17,17.25,8.17,'E3EAE9',.8)
# four biological scenes without surrounding boxes
for x,title in [(.85,'Diverse primer pairs'),(5.20,'Length-controlled design'),(9.55,'Fixed-primer completion'),(13.90,'Local redesign')]:
 text(x,8.5,3.48,.55,title,15,C['ink'],True)
# Scene 1: same DNA, alternative primers
for y,f,r in [(9.31,.12,2.17),(9.79,.49,2.50),(10.27,.30,1.9)]:duplex(.95,y,3.3,f,r,True)
text(.87,10.74,3.5,.29,'Multiple solutions for one target',11,C['muted'])
# Scene 2: varying target spans on same DNA
for y,f,r in [(9.31,.18,1.42),(9.79,.18,2.0),(10.27,.18,2.57)]:
 duplex(5.28,y,3.3,f,r,True)
text(5.21,10.74,3.5,.29,'Choose the amplicon range',11,C['muted'])
# Scene 3: fixed F or fixed R
for y,side in [(9.43,'F'),(10.15,'R')]:
 duplex(9.63,y,3.3,.22,2.43,True)
 lock(9.99 if side=='F' else 12.3,y-.32,C['teal'] if side=='F' else C['coral'])
 text(10.71,y-.3,1.24,.28,'Design R' if side=='F' else 'Design F',10,C['muted'],False,'center')
text(9.56,10.74,3.65,.29,'Keep either primer; complete its pair',10.5,C['muted'])
# Scene 4: old partner to new partner while anchor retained
duplex(13.99,9.43,3.2,.17,1.76,False);lock(14.29,9.1,C['teal'])
text(16.48,9.02,.8,.25,'Initial',10,C['muted'])
duplex(13.99,10.15,3.2,.17,2.5,True);lock(14.29,9.82,C['teal'])
arrow_curve([(16.02,9.76),(16.38,9.75),(16.64,9.90)],C['coral'],.035,.12)
text(13.91,10.74,3.6,.29,'Retain one primer; redesign its partner',10.5,C['muted'])
# Remove Office theme shadows for clean scientific rendering.
for shape in s.shapes:
 for el in shape._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):el.set('idx','0')
 etree.SubElement(shape._element.spPr,'{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
prs.save(OUT/'DiffuPrimer_method.pptx')
head='<svg xmlns="http://www.w3.org/2000/svg" width="1800" height="1130" viewBox="0 0 1800 1130"><defs><marker id="arrow" viewBox="0 0 10 10" refX="8" refY="5" markerWidth="5" markerHeight="5" orient="auto-start-reverse"><path d="M 0 0 L 10 5 L 0 10 z" fill="context-stroke"/></marker></defs>'
(OUT/'DiffuPrimer_method.svg').write_text(head+''.join(svg)+'</svg>')
print(f'Saved biological overview: {len(s.shapes)} native objects.')
