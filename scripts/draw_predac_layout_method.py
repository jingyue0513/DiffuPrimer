"""Fresh native PowerPoint method figure with a two-stage modular composition.
No imports or layout reuse from the previous method drawings.
"""
from pathlib import Path
import html
from pptx import Presentation
from pptx.util import Mm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from lxml import etree

OUT=Path('paper/figures/predac_layout_method'); OUT.mkdir(parents=True,exist_ok=True)
BASE='DiffuPrimer_method'
P=Presentation(); P.slide_width=Mm(178); P.slide_height=Mm(104)
S=P.slides.add_slide(P.slide_layouts[6]); V=[]
INK='29463E'; MUTED='62736D'; GREEN='77A38B'; BLUE='DCEAF1'; PALE='EDF4E7'; GOLD='F7EDCE'; F='308D78'; R='CF795D'; MASK='AA9ABB'

def box(x,y,w,h,fill='FFFFFF',stroke=None,r=.9):
 sh=S.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Mm(x),Mm(y),Mm(w),Mm(h))
 sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
 if stroke:sh.line.color.rgb=RGBColor.from_string(stroke);sh.line.width=Pt(.55)
 else:sh.line.fill.background()
 if r:sh.adjustments[0]=min(.13,r/min(w,h))
 V.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{r}" fill="#{fill}" stroke="#{stroke or fill}" stroke-width=".17"/>')
 return sh

def text(x,y,w,h,t,size=7,color=INK,bold=False,align='center',italic=False):
 sh=S.shapes.add_textbox(Mm(x),Mm(y),Mm(w),Mm(h));tf=sh.text_frame;tf.clear();tf.word_wrap=False
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,ln in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=ln;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.italic=italic;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(0);p.space_before=Pt(0)
  p.alignment={'center':PP_ALIGN.CENTER,'left':PP_ALIGN.LEFT,'right':PP_ALIGN.RIGHT}[align]
  xx=x+w/2 if align=='center' else x+w if align=='right' else x
  V.append(f'<text x="{xx}" y="{y+size*.35278*.9+i*size*.35278*1.18}" fill="#{color}" font-family="DejaVu Sans" font-size="{size*.35278}" font-weight="{700 if bold else 400}" font-style="{"italic" if italic else "normal"}" text-anchor="'+{'center':'middle','left':'start','right':'end'}[align]+f'">{html.escape(ln)}</text>')
 return sh

def line(x,y,xx,yy,color=GREEN,width=.65,arrow=False,dash=False):
 sh=S.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Mm(x),Mm(y),Mm(xx),Mm(yy));sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(width)
 if dash:
  from pptx.enum.dml import MSO_LINE_DASH_STYLE
  sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
 if arrow:
  e=etree.SubElement(sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln'),'{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd');e.set('type','triangle')
 V.append(f'<line x1="{x}" y1="{y}" x2="{xx}" y2="{yy}" stroke="#{color}" stroke-width="{width*.35278}"'+(' stroke-dasharray=".8,.6"' if dash else '')+(' marker-end="url(#a)"' if arrow else '')+'/>')

def arrow(x,y,xx,yy):line(x,y,xx,yy,arrow=True)
def down_tip(x,y):
 # Fixed-size triangle remains entirely inside the 1.1 mm layer gap.
 sh=S.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,Mm(x-.48),Mm(y),Mm(.96),Mm(.67));sh.rotation=180
 sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(GREEN);sh.line.fill.background()
 V.append(f'<polygon points="{x-.48},{y} {x+.48},{y} {x},{y+.67}" fill="#{GREEN}"/>')

def module(x,y,w,h,title):
 box(x,y,w,h,'FFFFFF',GREEN,3)
 text(x+1,y+1.6,w-2,7,title,8.1,bold=True)
def labelbox(x,y,w,h,t,fill=PALE,size=7):
 box(x,y,w,h,fill,'B7CBBE',.6);text(x+.4,y+.55,w-.8,h-1,t,size)
def matrix(x,y,rows,cols,cw,ch,colors):
 for row in range(rows):
  for col in range(cols):
   box(x+col*cw,y+row*ch,cw-.15,ch-.15,colors[(row*3+col*2)%len(colors)],None,0)
def states(x,y,st,cw=3,h=3.4):
 for i,c in enumerate(st):
  col={'0':'E8EDE9','F':F,'R':R,'M':MASK}[c]
  box(x+i*cw,y,cw-.1,h,col,None,0)
  text(x+i*cw,y+.35,cw-.1,h-.5,c,7,'FFFFFF' if c!='0' else MUTED,True)
def primer(x,y,w,right=True,color=F):line(x,y,x+w,y,color,1.8,arrow=right);line(x+w,y,x,y,color,1.8,arrow=not right) if not right else None
def pair(x,y,w,fs=2,rs=None):
 rs=w-6 if rs is None else rs
 box(x+fs,y-.5,rs-fs+4,2.6,'F7F0DB',None,0)
 line(x,y,x+w,y,'AFBFB6',.6);line(x,y+1.7,x+w,y+1.7,'AFBFB6',.6)
 primer(x+fs,y+1.7,4,True,F);primer(x+rs,y,4,False,R)

# Two separated horizontal stages, borrowing only the reference's composition.
box(.6,.6,176.8,56.4,'FFFFFF','B9C9A5',.4)
text(2,1.6,174,4,'Model construction',10,bold=True,italic=True)
module(2,7.5,29,47.5,'Reference data')
module(35,7.5,51,47.5,'Binding-map encoding')
module(90,7.5,41,47.5,'Conditional CNN')
module(135,7.5,41,47.5,'Denoiser training')
for a,b in [(31.4,34.5),(86.4,89.5),(131.4,134.5)]:arrow(a,31,b,31)

# Biological source data and paired labels.
labelbox(4.5,16,24,7,'RefSeq\ntranscripts',BLUE)
line(7,26.2,26,26.2,'789685',.8)
for a,b in [(8,10),(13,16),(21,24)]:box(a,24.8,b-a,2.7,'C9DDBC',None,0)
text(5,29,23,3,'ATCG template',7)
arrow(16.5,33,16.5,35.5)
labelbox(4.5,36.5,24,6,'Primer positions',PALE)
pair(5.5,46,22,2,15)
text(4,50,25,3,'Paired F / R sites',7,MUTED)

# Training corruption occupies one compact column; the matrix representation another.
text(37,15.8,29,3,'Fixed transcript',7,bold=True)
text(37,20.2,29,3,'A T G C T A G C A',7)
states(38,24.5,'0FFF00RR0',3)
arrow(51.5,28.8,51.5,31.8)
text(36.5,29.2,12,3,'t / T',7,MUTED)
states(38,33,'MFMFMMRM0',3)
text(37,37.8,29,3,'Noisy binding map',7)
# Illustrative tensors contain no invented numerical measurements.
text(67.4,16,16.8,3,'Features',7,bold=True)
matrix(69,20.1,6,5,2.7,2.2,['D9E9F0','BED6E4','A9C7D8','E8F0F4'])
text(67,34.2,18,3,'L × 96',7,MUTED)
labelbox(38,43.6,45,4.1,'9 design values + time',GOLD)
labelbox(38,49.1,45,4.1,'+ relative position',BLUE)

# Exact layer boxes; residual blocks are stated explicitly, not depicted as a generic network.
text(91.5,15.1,38,3,'Embed + add conditions',7)
labelbox(94,19.3,33,4.0,'96-channel input',BLUE)
ys=[25.0,30.0,35.0,40.0]
for i,(y,d) in enumerate(zip(ys,[1,4,16,64])):
 labelbox(94,y,33,3.9,f'Residual Conv1d · d = {d}',PALE)
 down_tip(110.5,y-.86)
down_tip(110.5,44.65)
labelbox(94,45.7,33,4.0,'1 × 1 head → 0 / F / R',BLUE)
text(92,51.2,37,3,'4 blocks · k = 5 · width 96',7,MUTED)

# Target supervision and learned output kept in a narrow compact module.
text(137,15.8,37,3,'Predict clean labels',7,bold=True)
states(139.5,20.4,'0FFF00RR0',3.15)
arrow(154,24.8,154,27.2)
labelbox(138,28.1,35,8.2,'Masked-site\ncross-entropy',GOLD)
text(138,38.2,35,3,'Clean x₀ target · T / t weight',7,MUTED)
arrow(154,42.7,154,45.0)
labelbox(138,46.1,35,6.8,'Trained denoiser',PALE,7.5)

box(.6,58.4,176.8,45,'FFFFFF','B9C9A5',.4)
text(2,59.4,174,4,'Conditional primer-pair design',10,bold=True,italic=True)
module(2,65.4,34,35.5,'Design request')
module(40,65.4,49,35.5,'Iterative recovery')
module(93,65.4,36,35.5,'Primer-pair output')
module(133,65.4,43,35.5,'Design modes')
for a,b in [(36.4,39.5),(89.4,92.5),(129.4,132.5)]:arrow(a,83.2,b,83.2)
labelbox(4.5,73.3,29,5,'Length · Tm · GC',GOLD)
labelbox(4.5,80,29,5,'Optional fixed F or R',BLUE)
states(5.5,88.3,'MMMMMMMMM',3)
text(4,95,30,3,'Unanchored: all M',7,MUTED)

# Short contained feedback loop: no cross-panel path or arrows through text.
labelbox(45,73.3,34,4.8,'Trained CNN → clean map',BLUE)
arrow(62,78.6,62,79.6)
labelbox(45,80.1,34,4.8,'Sample compatible F / R',PALE)
arrow(62,85.5,62,86.5)
labelbox(45,87,34,4.8,'Reveal M sites with 1 / t',GOLD)
line(79.7,89.4,84.5,89.4,dash=True);line(84.5,89.4,84.5,75.7,dash=True);line(84.5,75.7,79.7,75.7,arrow=True,dash=True)
text(42,93.4,45,3,'Eligibility · pair ΔTm · anchors',7,MUTED)
text(42,97,45,3,'T → 1; preserve revealed labels',7,MUTED)

pair(96,75.4,30,3,21)
text(95,80.3,32,3,'F: extract',7,color=F)
text(95,84.1,32,3,'R: reverse-complement',7,color=R)
labelbox(96,90.5,30,8.4,'Then assess\nfull-reference specificity',BLUE)

# Four compact biologically interpretable scenes instead of prose.
for xx,yy,title,mode in [(135,73.3,'Diversity',0),(155.5,73.3,'Length control',1),(135,87,'Fix F or R',2),(155.5,87,'Local redesign',3)]:
 text(xx,yy,18.5,3,title,7,bold=True)
 for q in range(2):
  fs=1.4+q*1.3 if mode==0 else 1.3
  rs=(11.3-q if mode==0 else 8.2+q*3.1 if mode in (1,3) else 11.2)
  pair(xx+.5,yy+5.1+q*3.4,17,fs,rs)
  if mode in (2,3):
   lockx=xx+2 if mode==3 or q==0 else xx+12
   box(lockx,yy+3.9+q*3.4,1.5,1.3,F if mode==3 or q==0 else R,None,.15)

for sh in S.shapes:
 for e in sh._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):e.set('idx','0')
 etree.SubElement(sh._element.spPr,'{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
P.save(OUT/f'{BASE}.pptx')
head='<svg xmlns="http://www.w3.org/2000/svg" width="178mm" height="104mm" viewBox="0 0 178 104"><defs><marker id="a" viewBox="0 0 10 10" refX="8" refY="5" markerWidth="5" markerHeight="5" orient="auto-start-reverse"><path d="M0 0 L10 5 L0 10Z" fill="context-stroke"/></marker></defs>'
(OUT/f'{BASE}.svg').write_text(head+''.join(V)+'</svg>')
print(f'{len(S.shapes)} native objects; 178 × 104 mm')
