"""Export and validate the macro/micro architecture and binding-position modeling diagram."""
from pathlib import Path
import base64, hashlib, json, re, shutil, subprocess
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
OUT=Path('paper/figures/architecture_method'); BASE='DiffuPrimer_method'
shutil.copyfile(OUT/'render_native'/f'{BASE}.pdf',OUT/f'{BASE}.pdf')
for name,args in [(BASE,['-r','350']),('preview',['-scale-to-x','1600','-scale-to-y','-1']),('print_review_700px',['-scale-to-x','700','-scale-to-y','-1'])]:
 subprocess.run(['pdftoppm','-png','-singlefile',*args,str(OUT/f'{BASE}.pdf'),str(OUT/name)],check=True)
with Image.open(OUT/f'{BASE}.png') as src:
 im=src.convert('RGB');pixels=im.size
 im.save(OUT/f'{BASE}.png',dpi=(350,350));im.save(OUT/f'{BASE}.tiff',dpi=(350,350),compression='tiff_lzw')
svgpath=OUT/f'{BASE}.svg';svg=svgpath.read_text();svg=re.sub('<style.*?</style>','',svg,flags=re.S)
css=[]
for fname,weight in [('DejaVuSans.ttf','400'),('DejaVuSans-Bold.ttf','700')]:
 data=base64.b64encode((Path('/usr/share/fonts/truetype/dejavu')/fname).read_bytes()).decode()
 css.append('@font-face{font-family:"DejaVu Sans";font-weight:'+weight+';src:url(data:font/ttf;base64,'+data+') format("truetype");}')
svgpath.write_text(svg.replace('<defs>','<defs><style>'+''.join(css)+'</style>',1))
prs=Presentation(OUT/f'{BASE}.pptx');shapes=list(prs.slides[0].shapes);sizes=[]
for sh in shapes:
 if sh.has_text_frame:
  for p in sh.text_frame.paragraphs:
   for run in p.runs:
    if run.text:sizes.append((run.font.size or p.font.size).pt)
assert min(sizes)>=7
assert not any(s.shape_type==MSO_SHAPE_TYPE.PICTURE for s in shapes)
page=PdfReader(OUT/f'{BASE}.pdf').pages[0]
assert abs(float(page.mediabox.width)*25.4/72-178)<.03
assert abs(float(page.mediabox.height)*25.4/72-116)<.03
sources=['src/diffuprimer/model.py','experiments/joint_tm/src/diffuprimer/decode.py','experiments/joint_tm/src/diffuprimer/sampling.py','experiments/joint_tm/src/diffuprimer/guidance.py','experiments/temperature_4/scripts/evaluate.py','experiments/final_followup/anchored_decode.py','scripts/draw_diffuprimer_architecture_method.py','scripts/export_diffuprimer_architecture_method.py','docs/references/predac_flu_b/reference_figure1.jpg','docs/references/dit_iccv2023/figure3.png','docs/references/convnext_cvpr2022/figure4.png']
report=dict(figure='DiffuPrimer original CNN: true macro/micro computational architecture with input-only additive conditioning and a separate fixed-template binding-map diffusion panel',dimensions_mm=[178,116],native_editable_objects=len(shapes),raster_objects_in_ppt=0,minimum_font_pt=min(sizes),raster_dpi=350,raster_pixels=pixels,reference=['PREDAC-FluB Figure 1, doi:10.1093/bib/bbaf308','DiT ICCV2023 Figure3: macro/detail computational graph representation','ConvNeXt CVPR2022 Figure4: residual identity branch representation'],reference_artwork_copied=False,source_hashes={p:hashlib.sha256(Path(p).read_bytes()).hexdigest() for p in sources},files={p.name:dict(bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in OUT.iterdir() if p.suffix in {'.pptx','.pdf','.svg','.png','.tiff'}},visual_review='Actual LibreOffice PowerPoint rendering and 700 px print preview inspected for text clipping, arrow placement and biological direction.')
(OUT/'provenance.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps({k:v for k,v in report.items() if k not in ('files','source_hashes')},indent=2))
