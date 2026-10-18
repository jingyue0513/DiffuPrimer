"""Editable PowerPoint and vector SVG method figure, from shared primitives."""
from pathlib import Path
import math, html
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from lxml import etree
OUT=Path('paper/figures/attention_style_draft'); OUT.mkdir(parents=True,exist_ok=True)
prs=Presentation(); prs.slide_width= Inches(18); prs.slide_height=Inches(10.8)
s=prs.slides.add_slide(prs.slide_layouts[6]); svg=[]
C={'ink':'193D50','muted':'607887','teal':'278E76','coral':'D97554','gold':'B78929','purple':'9185BD','gray':'DCE5E9','line':'CEDADF'}
def rect(x,y,w,h,fill='FFFFFF',stroke=None,r=.1):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
 sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(fill); sh.line.fill.background() if not stroke else None
 if stroke: sh.line.color.rgb=RGBColor.from_string(stroke); sh.line.width=Pt(.8)
 if r:
  try:sh.adjustments[0]=.08
  except Exception:pass
 svg.append(f'<rect x="{x*100}" y="{y*100}" width="{w*100}" height="{h*100}" rx="{r*100}" fill="#{fill}" stroke="#{stroke or fill}" stroke-width="1"/>')
 return sh
def text(x,y,w,h,t,size=16,color='193D50',bold=False,align='left'):
 sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 from pptx.enum.text import PP_ALIGN
 for i,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(0)
  p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
  xx=x*100+(w*50 if align=='center' else w*100 if align=='right' else 0)
  svg.append(f'<text x="{xx}" y="{y*100+size*100/72*.9+i*size*100/72*1.2}" font-family="DejaVu Sans" font-size="{size*100/72}" font-weight="{700 if bold else 400}" text-anchor="'+{'left':'start','center':'middle','right':'end'}[align]+f'" fill="#{color}">{html.escape(line)}</text>')
 return sh
def line(x1,y1,x2,y2,color='607887',width=1.5,arrow=False,dash=False):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2));sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(width)
 if arrow:
  el=etree.SubElement(sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln'),'{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd');el.set('type','triangle');el.set('w','med');el.set('len','med')
 svg.append(f'<line x1="{x1*100}" y1="{y1*100}" x2="{x2*100}" y2="{y2*100}" stroke="#{color}" stroke-width="{width*100/72}"'+(' marker-end="url(#arrow)"' if arrow else '')+'/>')
def cells(x,y,states,cell=.19):
 colors={'0':C['gray'],'F':C['teal'],'R':C['coral'],'M':C['purple']}
 for i,k in enumerate(states):
  rect(x+i*cell,y,cell-.025,.28,colors[k],r=.025)
  if cell>=.18:text(x+i*cell,y+.035,cell-.025,.2,k,9,'FFFFFF' if k!='0' else C['muted'],True,'center')
def pair(x,y,w=2.4,f=.35,r=1.8,muted=False):
 line(x,y,x+w,y,'B8C8CE',1.6); line(x+f,y-.10,x+f+.4,y-.1,C['teal'],3,True);line(x+r+.4,y+.1,x+r,y+.1,C['coral'],3,True)
def badge(x,y,n,title,w):
 rect(x,y,.36,.36,C['ink'],r=.08);text(x,y+.035,.36,.28,n,15,'FFFFFF',True,'center');text(x+.5,y+.015,w-.5,.45,title,19,C['ink'],True)
# Biological illustration: editable ribbons, nucleotide beads and primer glyphs.
def poly(points,fill):
 b=s.shapes.build_freeform(points[0][0]*1000,points[0][1]*1000,scale= Inches(1)/1000)
 b.add_line_segments([(x*1000,y*1000) for x,y in points[1:]],close=True)
 sh=b.convert_to_shape();sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background()
 svg.append('<polygon points="'+' '.join(f'{x*100},{y*100}' for x,y in points)+f'" fill="#{fill}"/>')
 return sh
def disk(x,y,r,fill,stroke=None):
 sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x-r), Inches(y-r), Inches(r*2), Inches(r*2));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
 if stroke:sh.line.color.rgb=RGBColor.from_string(stroke);sh.line.width=Pt(.8)
 else:sh.line.fill.background()
 svg.append(f'<circle cx="{x*100}" cy="{y*100}" r="{r*100}" fill="#{fill}" stroke="#{stroke or fill}"/>')
def ribbon(points,width,fill):
 upper=[];lower=[]
 for i,(x,y) in enumerate(points):
  a=points[max(0,i-1)];b=points[min(len(points)-1,i+1)];dx=b[0]-a[0];dy=b[1]-a[1];n=math.hypot(dx,dy) or 1
  upper.append((x-dy/n*width/2,y+dx/n*width/2));lower.append((x+dy/n*width/2,y-dx/n*width/2))
 poly(upper+lower[::-1],fill)
def helix(x,y,w,amp=.53,cycles=2.1,pale=False):
 col1,col2=('BFDAD8','D3CFE4') if pale else ('54A99E','9382B9')
 n=400
 for j in range(0,n+1,14):
  xx=x+w*j/n;a=amp*math.sin(j/n*math.pi*2*cycles)
  line(xx,y-a,xx,y+a,'E3DCCA' if pale else 'D5C5A0',1.9)
 for sign,col in [(1,col1),(-1,col2)]:
  pts=[(x+w*i/n,y+sign*amp*math.sin(i/n*math.pi*2*cycles)) for i in range(n+1)]
  ribbon(pts,.07 if pale else .065,col)
  # a narrower highlight gives the editable strand a soft ribbon appearance
  ribbon([(a,b-.016) for a,b in pts],.016,'E7F1EF' if pale else 'BDDDD7')
def arrow_curve(points,color,width=.08,head=.23):
 ribbon(points,width,color)
 x,y=points[-1];a,b=points[-2];dx=x-a;dy=y-b;n=math.hypot(dx,dy);ux=dx/n;uy=dy/n
 poly([(x+ux*head*.45,y+uy*head*.45),(x-ux*head+uy*head*.6,y-uy*head-ux*head*.6),(x-ux*head-uy*head*.6,y-uy*head+ux*head*.6)],color)
def ntrow(x,y,seq,step=.23,size=13):
 colors={'A':'448E84','T':'C98653','C':'6596BC','G':'A181B4'}
 for i,nt in enumerate(seq):
  disk(x+i*step+.1,y+.11,.105,'FFFFFF')
  text(x+i*step,y-.012,.20,.27,nt,size,colors[nt],True,'center')
def labels(x,y,states,step=.23):
 for i,k in enumerate(states):
  col={'F':C['teal'],'R':C['coral'],'M':'A395C4','0':'B7C8CC'}[k]
  if k=='0':disk(x+i*step+.10,y+.13,.037,col)
  else:
   disk(x+i*step+.10,y+.13,.107,col)
   text(x+i*step,y+.007,.20,.27,k,10,'FFFFFF',True,'center')
def primer(x,y,w,col,right=True):
 line(x,y,x+w,y,col,4)
 # tapered primer extension, not an administrative flowchart arrow
 end=x+w if right else x;dx=1 if right else -1
 poly([(end+dx*.12,y),(end-dx*.11,y-.09),(end-dx*.11,y+.09)],col)
def lock(x,y,col):
 pts=[(x+.075+.075*math.cos(math.pi-i*math.pi/20),y+.08-.09*math.sin(math.pi-i*math.pi/20)) for i in range(21)]
 ribbon(pts,.028,col);rect(x-.01,y+.07,.17,.14,col,r=.025);disk(x+.075,y+.125,.018,'FFFFFF')
def duplex(x,y,w,ff=.3,rr=2.4,shade=False):
 if shade:rect(x+ff,y-.015,rr-ff+.48,.20,'F1EBD9',r=.025)
 line(x,y,x+w,y,'BDCDCE',1.2);line(x,y+.17,x+w,y+.17,'BDCDCE',1.2)
 for i in range(int(w/.13)+1):line(x+i*.13,y+.03,x+i*.13,y+.14,'DFE5E4',.6)
 primer(x+ff,y+.13,.44,C['teal'],True);primer(x+rr,y+.035,.44,C['coral'],False)

def textline(x,y,w,t,size=17,color=None,bold=False,align='left'):
 return text(x,y,w,.42,t,size,color or C['ink'],bold,align)

def state_row(x,y,states):
 # Large, native editable characters; unselected positions remain explicit zeros.
 for i,k in enumerate(states):
  col={'F':C['teal'],'R':C['coral'],'M':C['purple'],'0':'B3C3C8'}[k]
  if k=='0':
   text(x+i*.24,y,.225,.30,'0',13,col,False,'center')
  else:
   disk(x+i*.24+.11,y+.155,.124,col)
   text(x+i*.24-.005,y+.02,.24,.28,k,12,'FFFFFF',True,'center')

def biological_state(x,heading,symbol,states):
 textline(x,1.41,4.5,heading,19,bold=True,align='center')
 textline(x,1.92,4.5,symbol,16,C['purple'],align='center')
 line(x-.04,2.9,x+4.24,2.9,'FFFFFF',22)
 ntrow(x,2.73,'ATGCGTACGATCGACTGA',.24,16)
 state_row(x,3.59,states)
 for i,k in enumerate(states):
  if k in 'FR':line(x+i*.24+.11,3.02,x+i*.24+.11,3.47,'B5D8CD' if k=='F' else 'E6BDAF',1)

rect(0,0,18,10.8,r=0)
text(.56,.18,17,.55,'DiffuPrimer: conditional diffusion for qPCR primer pairs',27,C['ink'],True)
textline(.59,.83,17,'a   Binding-position diffusion on a fixed transcript',18,bold=True)

# One faint molecular ribbon unifies the unchanged template across the mask states.
helix(.43,2.87,17.1,.54,4.8,True)
clean='00FFF000000RRR0000'
partial='0MFMF0MM0MMRMR0M0M'
biological_state(.80,'Clean primer-binding map','x₀',clean)
biological_state(6.85,'Partially masked labels','xₜ',partial)
biological_state(12.90,'Masked initialization','xT','M'*18)
for xa,xb in [(5.35,6.47),(11.40,12.51)]:
 arrow_curve([(xa,2.25),(xa+.28,2.19),(xb-.28,2.19),(xb,2.25)],'AA9CC9',.055,.15)
 textline(xa-.20,1.82,1.6,'+ mask',14,C['purple'],align='center')
 for j in range(5):disk(xa+.10+j*.185,3.72+.07*math.sin(j*2),.028+j%2*.01,'C5B7DD')
textline(.83,3.10,4.3,'ATCG template remains fixed',14,C['muted'],align='center')
textline(6.88,3.10,4.3,'Mask probability t / T',14,C['muted'],align='center')
textline(12.93,3.10,4.3,'Start from M at every position',14,C['muted'],align='center')

# Reverse process has the opposite direction and a continuous organic recovery path.
curve=[(15.04-12.12*i/150,4.17+.49*math.sin(math.pi*i/150)) for i in range(151)]
arrow_curve(curve,C['teal'],.072,.21)
textline(5.26,4.12,7.4,'Learn conditional recovery; sample t = T → 1',17,C['teal'],True,'center')
primer(1.43,4.63,.64,C['teal'],True);primer(3.55,4.63,.64,C['coral'],False)
textline(.67,4.94,4.80,'F = template span; R = reverse complement',13,C['muted'],align='center')

# Soft band borrows PREDAC-FluB's train/inference hierarchy, without module boxes.
rect(.40,5.46,17.20,2.15,'F5F9F3',r=.10)
textline(.69,5.64,16.5,'b   Conditional attention and structured recovery at each step',18,bold=True)

# Conditions: biological ruler, fixed oligo and compact sequence-aware conditions.
textline(.83,6.19,4.2,'Template + design conditions',16,bold=True)
line(.98,6.89,2.20,6.89,'B58F43',2)
for i in range(8):line(.98+i*.174,6.81,.98+i*.174,6.98,'B58F43',1)
lock(2.66,6.75,C['teal']);primer(3.03,6.89,.65,C['teal'])
textline(.82,7.18,4.45,'Length · Tm · GC · optional fixed F / R',13,C['muted'])
arrow_curve([(4.55,6.92),(4.77,6.92),(4.98,6.92)],'A9B09A',.035,.11)

# Editable multi-head attention glyph: token interactions across distant F / R sites.
textline(5.17,6.10,6.18,'Conditional Transformer denoiser',17,bold=True)
tokenx=[5.35+i*.41 for i in range(8)]
for a,b,col,amp in [(0,7,C['teal'],.54),(1,6,'9787B9',.41),(2,7,'C59A4C',.30),(0,4,'D09179',.23)]:
 pts=[]
 for i in range(45):
  q=i/44;pts.append((tokenx[a]+(tokenx[b]-tokenx[a])*q,7.03-amp*math.sin(math.pi*q)))
 ribbon(pts,.012,col)
for i,(nx,k) in enumerate(zip(tokenx,'FM00MMRM')):
 disk(nx,7.07,.085,C['teal'] if k=='F' else C['coral'] if k=='R' else C['purple'] if k=='M' else 'B8C8C7')
 text(nx-.09,6.96,.18,.23,k,9,'FFFFFF',True,'center')
textline(8.79,6.65,2.61,'Predict 0 / F / R',14,C['muted'])
textline(5.16,7.18,6.46,'4 attention blocks · rotary positions · condition/time modulation',12.2,C['muted'])
arrow_curve([(11.47,6.92),(11.67,6.92),(11.89,6.92)],'A9B09A',.035,.11)

textline(12.10,6.19,5.1,'Compatible primer-pair sampling',16,bold=True)
duplex(12.14,6.78,4.49,.31,3.65,True)
textline(12.10,7.15,5.0,'Lengths · span order · anchors · joint ΔTm',12.2,C['muted'])
# Upward link identifies structured sampling as part of reverse diffusion.
arrow_curve([(15.15,5.99),(15.36,5.66),(15.42,5.10),(15.04,4.62)],'91B9A2',.028,.13)
textline(12.23,4.93,5.18,'Reveal with 1 / t; preserve fixed labels',13,C['muted'],align='center')

# Four biological task scenes use invariant F/R colors and consistent template glyphs.
textline(.61,7.99,16.9,'c   One model, four flexible design modes',18,bold=True)
positions=[.70,5.10,9.49,13.89]
for x,title in zip(positions,['Diverse primer pairs','Length control','Fixed-primer completion','Local redesign']):
 textline(x,8.51,3.6,title,17,bold=True)
for y,f,r in [(9.20,.13,2.16),(9.64,.42,2.45),(10.08,.24,1.85)]:duplex(.79,y,3.45,f,r,True)
for y,r in [(9.20,1.49),(9.64,2.06),(10.08,2.64)]:duplex(5.19,y,3.45,.19,r,True)
for y,side in [(9.20,'F'),(10.00,'R')]:
 duplex(9.58,y,3.45,.20,2.50,True)
 lock(9.93 if side=='F' else 12.25,y-.29,C['teal'] if side=='F' else C['coral'])
 textline(10.55,y-.30,1.53,'Design R' if side=='F' else 'Design F',12,C['muted'],align='center')
duplex(13.98,9.20,3.36,.20,1.78,False);lock(14.33,8.91,C['teal'])
duplex(13.98,10.00,3.36,.20,2.50,True);lock(14.33,9.71,C['teal'])
arrow_curve([(16.09,9.58),(16.36,9.57),(16.63,9.75)],C['coral'],.035,.12)
textline(.70,10.49,3.65,'Alternative F / R positions',12.2,C['muted'])
textline(5.10,10.49,3.65,'Specified amplicon range',12.2,C['muted'])
textline(9.49,10.49,3.78,'Keep either F or R',12.2,C['muted'])
textline(13.89,10.49,3.65,'Change the partner position',12.2,C['muted'])

# Remove Office theme shadows; all graphics remain native PowerPoint geometry.
for shape in s.shapes:
 for el in shape._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):el.set('idx','0')
 etree.SubElement(shape._element.spPr,'{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
prs.save(OUT/'DiffuPrimer_attention_method.pptx')
head='<svg xmlns="http://www.w3.org/2000/svg" width="1800" height="1080" viewBox="0 0 1800 1080"><defs><marker id="arrow" viewBox="0 0 10 10" refX="8" refY="5" markerWidth="5" markerHeight="5" orient="auto-start-reverse"><path d="M 0 0 L 10 5 L 0 10 z" fill="context-stroke"/></marker></defs>'
(OUT/'DiffuPrimer_attention_method.svg').write_text(head+''.join(svg)+'</svg>')
print(f'Saved attention method draft: {len(s.shapes)} native editable objects.')
